# -*- coding: utf-8 -*-
import openpyxl
import json
import re
from datetime import datetime, date, time
from pathlib import Path

path = Path(r"c:\Users\teste\Downloads\Credenciamento de engenheiros atividades e municipios- Final (3).xlsx")
out = Path(r"c:\Users\teste\Desktop\sorteador")
wb = openpyxl.load_workbook(path, data_only=True)
ws = wb["Credenciamento 2020"]

ATIVIDADES = {
    "A": "Laudo Simplificado",
    "B": "Laudo Completo",
    "C": "Atividade C",
    "D": "Atividade D",
    "E": "Atividade E",
    "F": "Atividade F",
}
REGIOES = {
    "I": "Central (MG)",
    "II": "Zona da Mata (MG)",
    "III": "Sul de Minas",
    "IV": "Triângulo (MG)",
    "V": "Alto Paranaíba (MG)",
    "VI": "Centro Oeste (MG)",
    "VII": "Noroeste de Minas",
    "VIII": "Norte de Minas",
    "IX": "Jequitinhonha/Mucuri",
    "X": "Rio Doce (MG)",
    "XI": "Acre",
    "XII": "Amazonas",
    "XIII": "Rondônia",
    "XIV": "Amapá",
    "XV": "Pará",
    "XVI": "Maranhão",
    "XVII": "Piauí",
    "XVIII": "Ceará",
    "XIX": "Rio Grande do Norte",
    "XX": "Paraíba",
    "XXI": "Pernambuco",
    "XXII": "Alagoas",
    "XXIII": "Sergipe",
    "XXIV": "Bahia",
    "XXV": "Espírito Santo",
    "XXVI": "Rio de Janeiro",
    "XXVII": "São Paulo",
    "XXVIII": "Paraná",
    "XXIX": "Santa Catarina",
    "XXX": "Rio Grande do Sul",
    "XXXI": "Mato Grosso do Sul",
    "XXXII": "Mato Grosso",
    "XXXIII": "Goiás/Distrito Federal",
    "XXXIV": "Tocantins",
    "XXXV": "Roraima",
}
CODES_AT = ["A", "B", "C", "D", "E", "F"]
CODES_MUN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]
CODES_EST = [
    "XI", "XII", "XIII", "XIV", "XV", "XVI", "XVII", "XVIII", "XIX", "XX",
    "XXI", "XXII", "XXIII", "XXIV", "XXV", "XXVI", "XXVII", "XXVIII", "XXIX",
    "XXX", "XXXI", "XXXII", "XXXIII", "XXXIV", "XXXV",
]


def clean(s):
    if s is None:
        return ""
    if isinstance(s, datetime):
        return s.strftime("%Y-%m-%d")
    if isinstance(s, date):
        return s.isoformat()
    if isinstance(s, time):
        return s.strftime("%H:%M")
    return str(s).strip()


def marked(v):
    if v is None:
        return False
    return str(v).strip().upper() in ("X", "1", "S", "SIM")


def php_str(s):
    return "'" + str(s).replace("\\", "\\\\").replace("'", "\\'") + "'"


def parse_cidade_uf(contato):
    c = contato.replace("\xa0", " ")
    if re.search(r"\bBH\s*/\s*MG\b", c, re.I):
        return "Belo Horizonte", "MG"
    m = re.search(
        r"([A-Za-zÀ-ú.\s\-]+?)\s*/\s*(MG|SP|RJ|ES|DF|GO|BA|PR|SC|RS|MT|MS|PA|AM|CE|PE|PB|RN|AL|SE|MA|PI|TO|RO|AC|AP|RR)\b",
        c,
        re.I,
    )
    if m:
        cidade = re.sub(r"\s+", " ", m.group(1)).strip(" ,.-")
        if len(cidade) > 40:
            cidade = cidade.split(",")[-1].strip()
        if cidade.upper() in ("BH", "B.H.", "B. H."):
            cidade = "Belo Horizonte"
        return cidade, m.group(2).upper()
    if re.search(r"Belo Horizonte", c, re.I):
        return "Belo Horizonte", "MG"
    if re.search(r"Uberlândia|Uberlandia", c, re.I):
        return "Uberlândia", "MG"
    if re.search(r"Juiz de Fora", c, re.I):
        return "Juiz de Fora", "MG"
    return "Belo Horizonte", "MG"


companies = []
idx = 1
for row in ws.iter_rows(min_row=5, max_row=ws.max_row, values_only=True):
    nome = clean(row[5])
    if not nome:
        continue

    status = "ativa"
    if "não executam" in nome.lower() or "nao executam" in nome.lower():
        status = "inativa"
        nome = re.sub(r"\s*-\s*N[aã]o executam.*$", "", nome, flags=re.I).strip()

    contato = clean(row[2])
    lines = [l.strip() for l in contato.replace("\r", "").split("\n") if l.strip()]
    endereco = " ".join(lines[1:]) if len(lines) > 1 else (lines[0] if lines else "")
    endereco = re.sub(r"\s+", " ", endereco).strip()
    cidade, uf = parse_cidade_uf(contato)
    cep_m = re.search(r"CEP\s*([\d.\-]+)", contato, re.I)
    cep = cep_m.group(1) if cep_m else ""

    at_codes = [c for i, c in enumerate(CODES_AT) if marked(row[6 + i])]
    mun_codes = [c for i, c in enumerate(CODES_MUN) if marked(row[12 + i])]
    est_codes = [c for i, c in enumerate(CODES_EST) if marked(row[22 + i])]

    email = clean(row[57]) if len(row) > 57 else ""
    if ";" in email:
        email = email.split(";")[0].strip()
    telefone = clean(row[58]) if len(row) > 58 else ""
    responsavel = ""
    rm = re.search(r"\(([^0-9)]{3,})\)\s*$", telefone)
    if rm:
        responsavel = rm.group(1).strip()

    data_hab = clean(row[0])
    contrato = clean(row[4])
    ordem = row[3]

    companies.append(
        {
            "id": f"EMP-{idx:03d}",
            "nome": nome,
            "cnpj": "",
            "cidade": cidade,
            "uf": uf,
            "endereco": endereco[:220] if endereco else "Endereço conforme credenciamento",
            "bairro": "",
            "cep": cep,
            "servicos": [ATIVIDADES[c] for c in at_codes] or ["Credenciamento"],
            "atividadesCodigos": at_codes,
            "regioesMunicipios": [f"{c} - {REGIOES[c]}" for c in mun_codes],
            "regioesMunicipiosCodigos": mun_codes,
            "estadosAtuacao": [f"{c} - {REGIOES[c]}" for c in est_codes],
            "estadosAtuacaoCodigos": est_codes,
            "status": status,
            "dataConvenio": data_hab if data_hab else "2020-05-01",
            "numeroContrato": contrato,
            "ordemCredenciamento": int(ordem) if isinstance(ordem, (int, float)) else idx,
            "telefone": telefone,
            "email": email,
            "responsavel": responsavel,
            "observacoes": "Importado da planilha Credenciamento 2020",
            "latitude": None,
            "longitude": None,
        }
    )
    idx += 1

(out / "storage").mkdir(exist_ok=True)
json_path = out / "storage" / "empresas.json"
json_path.write_text(json.dumps(companies, ensure_ascii=False, indent=2), encoding="utf-8")
print("Wrote", json_path, "count=", len(companies))

lines = [
    "<?php",
    "",
    "declare(strict_types=1);",
    "",
    "/**",
    " * Empresas importadas da planilha Credenciamento 2020.",
    " * Persistência adicional em storage/empresas.json",
    " */",
    "",
    "return [",
]
for e in companies:
    lines.append("    [")
    for k, v in e.items():
        if v is None:
            lines.append(f"        '{k}' => null,")
        elif isinstance(v, bool):
            lines.append(f"        '{k}' => {'true' if v else 'false'},")
        elif isinstance(v, (int, float)):
            val = int(v) if float(v) == int(v) else v
            lines.append(f"        '{k}' => {val},")
        elif isinstance(v, list):
            items = ", ".join(php_str(x) for x in v)
            lines.append(f"        '{k}' => [{items}],")
        else:
            lines.append(f"        '{k}' => {php_str(v)},")
    lines.append("    ],")
lines.append("];")
lines.append("")
(out / "mocks" / "empresas.php").write_text("\n".join(lines), encoding="utf-8")
print("Wrote mocks/empresas.php")

php_cat = [
    "<?php",
    "",
    "declare(strict_types=1);",
    "",
    "return [",
    "    'atividades' => [",
]
php_cat += [f"        '{k}' => {php_str(v)}," for k, v in ATIVIDADES.items()]
php_cat += ["    ],", "    'regioesMunicipios' => ["]
php_cat += [f"        '{k}' => {php_str(REGIOES[k])}," for k in CODES_MUN]
php_cat += ["    ],", "    'estadosAtuacao' => ["]
php_cat += [f"        '{k}' => {php_str(REGIOES[k])}," for k in CODES_EST]
php_cat += ["    ],", "];", ""]
(out / "mocks" / "catalogo_credenciamento.php").write_text("\n".join(php_cat), encoding="utf-8")
print("cities:", sorted(set(c["cidade"] for c in companies))[:20])
print("total:", len(companies))
